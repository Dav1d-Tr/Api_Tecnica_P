"""FastAPI para planos de Demetalicos, por clase y tipo.

Conserva los endpoints historicos de cotizaciones PC.

"""

from __future__ import annotations

import hmac

import os

import re

import time

import unicodedata

from collections import deque

from dataclasses import dataclass

from pathlib import Path

from fastapi import Depends, FastAPI, HTTPException, Query, Security

from fastapi.responses import FileResponse, PlainTextResponse

from fastapi.security import APIKeyHeader

from pypdf import PdfReader

# Raices confirmadas en el servidor Demetalicos. Pueden reemplazarse por

# rutas UNC si FastAPI corre como servicio y no ve las unidades mapeadas.

RAICES = {

    "TECNICA": Path(os.environ.get("PLANOS_RAIZ_TECNICA", r"P:\TECNICA")),

    "TECNICA_2": Path(os.environ.get("PLANOS_RAIZ_TECNICA_2", r"P:\aplic\TECNICA 2")),

    "FS_DPTO": Path(os.environ.get("PLANOS_RAIZ_FS_DPTO", r"S:\FS-DPTO")),

}

# Ubicacion PC confirmada: P:\TECNICA\PLANOS DE COTIZACION.

RUTA_PC = Path(os.environ.get(

    "PLANOS_RUTA_PC", str(RAICES["TECNICA"] / "PLANOS DE COTIZACI\u00d3N")

))

API_KEY = os.environ.get("API_KEY", "")

PUBLIC_URL = os.environ.get("PUBLIC_URL", "").rstrip("/")

MAX_ELEMENTOS = int(os.environ.get("PLANOS_MAX_ELEMENTOS", "50000"))

MAX_SEGUNDOS = int(os.environ.get("PLANOS_MAX_SEGUNDOS", "15"))

MAX_PROFUNDIDAD = int(os.environ.get("PLANOS_MAX_PROFUNDIDAD", "14"))

MAX_PDF_MB = int(os.environ.get("PLANOS_MAX_PDF_MB", "50"))

if len(API_KEY) < 20:

    raise RuntimeError("Define API_KEY con al menos 20 caracteres.")

@dataclass(frozen=True)

class TipoPlano:

    clase: str

    tipo: str

    ruta: Path

def _ruta(raiz: str, *partes: str) -> Path:

    return RAICES[raiz].joinpath(*partes)

# Equivalencias PROVISIONALES de la imagen antigua (no confirmadas por usuario):

# P: de la imagen -> raiz TECNICA; H: -> TECNICA_2; S: -> FS_DPTO.

# Solo la ruta PC esta expresamente confirmada. Ajustar catalogo si difiere.

CATALOGO: tuple[TipoPlano, ...] = (

    TipoPlano("De todo", "PC", RUTA_PC),

    TipoPlano("Estanter\u00eda", "EP", _ruta("TECNICA", "BIBLIOTECA DE PRODUCTOS DM", "ESTANTERIA", "ESTANTERIA PERNADA", "PLANOS ESTANTERIA PERNADA")),

    TipoPlano("Estanter\u00eda", "ES", _ruta("TECNICA_2", "Dibujo", "0Planos", "ESTANTERIA PESADA")),

    TipoPlano("Estanter\u00eda", "HD", _ruta("TECNICA_2", "HOMECENTER")),

    TipoPlano("Estanter\u00eda", "HE", _ruta("TECNICA_2", "HOMECENTER", "1HOMECENTER ESTANTERIA")),

    TipoPlano("Estanter\u00eda", "CA", _ruta("FS_DPTO", "Proyectos", "Biblioteca de Productos", "Productos de Linea", "Estanteria", "Accesorios")),

    TipoPlano("G\u00f3ndola", "DG", _ruta("TECNICA", "GONDOLA")),

    TipoPlano("Por cliente", "SL", _ruta("TECNICA_2", "Dibujo", "0Planos", "SPORTLINE")),

    TipoPlano("Por cliente", "EX", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "EXW", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "ARR", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "CAR", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "CD", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "DM", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "JYB", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "OXX", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "TRU", RAICES["TECNICA_2"]),

    TipoPlano("Por cliente", "SHI", RAICES["TECNICA_2"]),

    TipoPlano("Kioscos", "KV", _ruta("TECNICA_2", "Dibujo", "0Planos")),

    TipoPlano("Kioscos", "KD", _ruta("TECNICA_2", "Dibujo", "0Planos")),

    TipoPlano("Exhibici\u00f3n", "DM", _ruta("TECNICA_2", "Dibujo", "0Planos", "DEMETALI")),

    TipoPlano("Exhibici\u00f3n", "CC", _ruta("TECNICA_2", "Dibujo", "0Planos", "DEMETALI")),

    TipoPlano("Exhibici\u00f3n", "BR", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "FL", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "GA", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "P", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "T", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "VE", _ruta("TECNICA_2", "Dibujo", "0Planos", "MISCELANEA")),

    TipoPlano("Exhibici\u00f3n", "OT", _ruta("TECNICA_2", "Dibujo", "0Planos", "OTROS", "otros1")),

)

app = FastAPI(

    title="API de Planos Demetalicos",

    description="Busca, lista, lee y descarga PDF de planos segun clase y tipo.",

    version="3.1.0",

    servers=[{"url": PUBLIC_URL}] if PUBLIC_URL else None,

)

esquema_header = APIKeyHeader(name="X-API-Key", auto_error=False)

def verificar_clave(por_header: str | None = Security(esquema_header)):

    if not hmac.compare_digest((por_header or "").encode(), API_KEY.encode()):

        raise HTTPException(status_code=401, detail="API key invalida o ausente.")

def _normalizar(texto: str) -> str:

    """Normaliza acentos y mayusculas solo para comparar textos."""

    texto = unicodedata.normalize("NFKD", texto)

    return "".join(c for c in texto if not unicodedata.combining(c)).casefold().strip()

def seleccionar_tipo(tipo: str, clase: str | None = None) -> TipoPlano:

    tipo = tipo.strip().upper()

    candidatos = [t for t in CATALOGO if t.tipo == tipo]

    if not candidatos:

        raise HTTPException(status_code=400, detail=f"Tipo '{tipo}' no definido. Consulta /api/planos/tipos.")

    if clase:

        candidatos = [t for t in candidatos if _normalizar(t.clase) == _normalizar(clase)]

        if not candidatos:

            raise HTTPException(status_code=400, detail=f"El tipo '{tipo}' no pertenece a la clase '{clase}'.")

    if len(candidatos) != 1:

        clases = ", ".join(t.clase for t in candidatos)

        raise HTTPException(status_code=400, detail=f"El tipo '{tipo}' es ambiguo. Indica clase: {clases}.")

    return candidatos[0]

def carpeta_valida(ruta: Path) -> Path:

    try:

        if not ruta.is_dir():

            raise HTTPException(status_code=503, detail="La carpeta de planos no esta disponible en el servidor.")

        return ruta

    except OSError as exc:

        raise HTTPException(status_code=503, detail=f"Error accediendo a la ruta configurada: {exc}") from exc

def ruta_segura(base: Path, relativa: str) -> Path:

    """Restringe cualquier ruta de entrada al arbol autorizado por el catalogo."""

    relativa = relativa.strip().replace("\\", "/")

    if not relativa:

        return base.resolve()

    partes = relativa.split("/")

    if (

        relativa.startswith("/")

        or ":" in relativa

        or any(p in {"", ".", ".."} for p in partes)

        or "\x00" in relativa

    ):

        raise HTTPException(status_code=400, detail="Ruta relativa no valida.")

    raiz = base.resolve()

    destino = raiz.joinpath(*partes).resolve()

    if destino != raiz and raiz not in destino.parents:

        raise HTTPException(status_code=403, detail="Ruta fuera de la carpeta autorizada.")

    return destino

def normalizar_pc(codigo_pc: str) -> str:

    codigo = codigo_pc.strip().upper()

    if codigo.isdecimal():

        codigo = "PC" + codigo

    if not re.fullmatch(r"PC\d+", codigo):

        raise HTTPException(status_code=400, detail="El codigo debe tener formato PC14774 o 14774.")

    return codigo

def carpeta_pc(codigo_pc: str) -> Path:

    raiz = carpeta_valida(seleccionar_tipo("PC").ruta)

    codigo = normalizar_pc(codigo_pc)

    encontrados = []

    try:

        for entrada in raiz.iterdir():

            if not entrada.is_dir():

                continue

            nombre = entrada.name.strip().upper()

            if nombre == codigo or nombre.startswith((codigo + " ", codigo + "-")):

                encontrados.append(entrada)

    except OSError as exc:

        raise HTTPException(status_code=503, detail=f"No se puede consultar cotizaciones: {exc}") from exc

    if not encontrados:

        raise HTTPException(status_code=404, detail=f"No se encontro carpeta para {codigo}.")

    encontrados.sort(key=lambda p: p.name.casefold())

    return encontrados[0]

def carpeta_planos_pc(codigo_pc: str) -> Path:

    cotizacion = carpeta_pc(codigo_pc)

    try:

        for entrada in cotizacion.iterdir():

            if entrada.is_dir() and entrada.name.strip().casefold() == "01 - planos":

                return entrada

    except OSError as exc:

        raise HTTPException(status_code=503, detail=f"Error abriendo carpeta PC: {exc}") from exc

    raise HTTPException(status_code=404, detail=f"No existe '01 - Planos' en '{cotizacion.name}'.")

def explorar_pdfs(base: Path, inicio: Path, consulta: str, recursivo: bool, limite: int) -> dict:

    """Busqueda acotada para no bloquear una unidad de red extensa."""

    carpeta_valida(inicio)

    cola = deque([(inicio, 0)])

    pdfs = []

    explorados = 0

    parcial = False

    motivos = set()

    fecha_inicio = time.monotonic()

    termino = _normalizar(consulta) if consulta else ""

    detener = False

    while cola and not detener:

        actual, profundidad = cola.popleft()

        try:

            with os.scandir(actual) as entradas:

                for entrada in entradas:

                    if time.monotonic() - fecha_inicio >= MAX_SEGUNDOS:

                        parcial, detener = True, True

                        motivos.add("tiempo maximo de busqueda")

                        break

                    explorados += 1

                    if explorados > MAX_ELEMENTOS:

                        parcial, detener = True, True

                        motivos.add("limite de elementos explorados")

                        break

                    try:

                        if entrada.is_symlink():

                            continue

                        if entrada.is_dir(follow_symlinks=False):

                            if recursivo:

                                if profundidad < MAX_PROFUNDIDAD:

                                    cola.append((Path(entrada.path), profundidad + 1))

                                else:

                                    parcial = True

                                    motivos.add("profundidad maxima")

                            continue

                        if not entrada.is_file(follow_symlinks=False) or not entrada.name.casefold().endswith(".pdf"):

                            continue

                        archivo = Path(entrada.path)

                        relativa = archivo.relative_to(base).as_posix()

                        if termino and termino not in _normalizar(relativa):

                            continue

                        pdfs.append({

                            "nombre": archivo.name,

                            "ruta_en_planos": relativa,

                            "tamano_bytes": entrada.stat(follow_symlinks=False).st_size,

                        })

                        if len(pdfs) >= limite:

                            parcial, detener = True, True

                            motivos.add("limite de PDF devueltos")

                            break

                    except OSError:

                        parcial = True

                        motivos.add("elementos sin acceso")

        except OSError:

            parcial = True

            motivos.add("carpetas sin acceso")

    pdfs.sort(key=lambda x: x["ruta_en_planos"].casefold())

    return {

        "total_pdfs": len(pdfs),

        "pdfs": pdfs,

        "elementos_explorados": explorados,

        "parcial": parcial,

        "motivo_parcial": ", ".join(sorted(motivos)) if motivos else None,

    }

def resumen_ruta_compartida(categoria: TipoPlano) -> str | None:

    iguales = [x for x in CATALOGO if x.ruta == categoria.ruta]

    if len(iguales) < 2:

        return None

    return (

        "Esta carpeta es compartida por varios tipos. La API busca por el texto "

        "de la consulta en nombres de PDF y subcarpetas; no puede separar "

        "tipos sin una convencion de nombres. Usa codigos completos."

    )

def localizar_pdf(tipo: str, clase: str | None, archivo: str, codigo_pc: str | None = None):

    categoria = seleccionar_tipo(tipo, clase)

    if categoria.tipo == "PC":

        if not codigo_pc:

            raise HTTPException(status_code=400, detail="codigo_pc es obligatorio para tipo PC.")

        base = carpeta_planos_pc(codigo_pc)

    else:

        base = carpeta_valida(categoria.ruta)

    if not archivo.strip().casefold().endswith(".pdf"):

        raise HTTPException(status_code=400, detail="El archivo debe ser un PDF.")

    destino = ruta_segura(base, archivo)

    if destino.is_file():

        return categoria, destino

    # Compatibilidad PC: se admitia pasar solo el nombre de un PDF anidado.

    if categoria.tipo == "PC" and "/" not in archivo and "\\\\" not in archivo:

        coincidencias = [

            p for p in base.rglob("*")

            if p.is_file() and p.suffix.casefold() == ".pdf" and p.name.casefold() == archivo.casefold()

            and base.resolve() in p.resolve().parents

        ]

        if len(coincidencias) == 1:

            return categoria, coincidencias[0]

        if len(coincidencias) > 1:

            raise HTTPException(

                status_code=409,

                detail="Hay varios PDF con ese nombre. Usa exactamente 'ruta_en_planos' devuelta por la busqueda.",

            )

    raise HTTPException(status_code=404, detail="PDF no encontrado. Usa 'ruta_en_planos' del resultado de busqueda.")

def obtener_texto(categoria: TipoPlano, pdf: Path, max_caracteres: int, codigo_pc: str | None = None):

    try:

        tamano_maximo = MAX_PDF_MB * 1024 * 1024

        if pdf.stat().st_size > tamano_maximo:

            raise HTTPException(status_code=413, detail=f"PDF supera el limite de {MAX_PDF_MB} MB.")

        lector = PdfReader(str(pdf))

        total_paginas = len(lector.pages)

        secciones = []

        caracteres = 0

        truncado = False

        for n, pagina in enumerate(lector.pages, start=1):

            try:

                texto_pagina = (pagina.extract_text() or "").strip()

            except Exception:

                texto_pagina = ""

            if texto_pagina:

                seccion = f"[Pagina {n}]\n{texto_pagina}"

                secciones.append(seccion)

                caracteres += len(seccion) + 2

            if caracteres > max_caracteres:

                truncado = True

                break

        texto_final = "\n\n".join(secciones)

    except HTTPException:

        raise

    except Exception as exc:

        raise HTTPException(status_code=500, detail=f"No se pudo leer el PDF: {exc}") from exc

    resultado = {

        "tipo": categoria.tipo,

        "clase": categoria.clase,

        "pdf": pdf.name,

        "paginas": total_paginas,

        "texto_extraido": bool(texto_final.strip()),

        "truncado": truncado or len(texto_final) > max_caracteres,

        "texto": texto_final[:max_caracteres],

        "observacion": None if texto_final.strip() else "Sin texto extraible. Podria ser un PDF escaneado.",

    }

    if codigo_pc:

        resultado["codigo_pc"] = normalizar_pc(codigo_pc)

    return resultado

@app.get("/", include_in_schema=False)

def inicio():

    return {"mensaje": "API de Planos funcionando", "documentacion": "/docs"}

@app.get("/robots.txt", response_class=PlainTextResponse, include_in_schema=False)

def robots():

    return "User-agent: *\nDisallow: /\n"

@app.get("/api/ping", summary="Conectividad publica", operation_id="pingApi")

def ping():

    return {"ok": True, "mensaje": "FastAPI esta funcionando"}

@app.get("/api/planos/estado", dependencies=[Depends(verificar_clave)], operation_id="estadoApi")

def estado():

    return {

        "accesible": RUTA_PC.is_dir(),  # Campo original para compatibilidad.

        "raices": {nombre: raiz.is_dir() for nombre, raiz in RAICES.items()},

        "rutas_configuradas": len(CATALOGO),

    }

@app.get("/api/planos/tipos", dependencies=[Depends(verificar_clave)], operation_id="listarTiposPlanos")

def tipos_planos():

    return {

        "total": len(CATALOGO),

        "tipos": [

            {"clase": t.clase, "tipo": t.tipo, "ruta": str(t.ruta), "accesible": t.ruta.is_dir()}

            for t in CATALOGO

        ],

    }

@app.get("/api/planos/buscar", dependencies=[Depends(verificar_clave)], operation_id="buscarPlanos")

def buscar_planos(

    codigo_pc: str | None = Query(None, description="PC14774 o 14774 (compatibilidad PC)"),

    recursivo: bool | None = Query(None, description="Por defecto: False en PC; True en otros tipos"),

    tipo: str = Query("PC", description="PC, EP, ES, HD, HE, CA, DG, SL, etc."),

    clase: str | None = Query(None, description="Requerida si tipo=DM"),

    consulta: str | None = Query(None, description="Codigo o parte del nombre de PDF o carpeta para tipos distintos de PC"),

    limite: int | None = Query(None, ge=1, le=5000),

):

    categoria = seleccionar_tipo(tipo, clase)

    if categoria.tipo == "PC":

        codigo = normalizar_pc(codigo_pc or consulta or "")

        carpeta = carpeta_planos_pc(codigo)

        listado = explorar_pdfs(carpeta, carpeta, "", bool(recursivo), limite or 5000)

        # Campo de salida historico: ruta respecto de la carpeta principal de PC.

        for pdf in listado["pdfs"]:

            pdf["ruta_relativa"] = (Path(carpeta.parent.name) / carpeta.name / pdf["ruta_en_planos"]).as_posix()

        return {

            "clase": categoria.clase,

            "tipo": "PC",

            "codigo_pc": codigo,

            "carpeta_pc": carpeta.parent.name,

            "carpeta_planos": carpeta.name,

            **listado,

        }

    termino = (consulta or codigo_pc or "").strip()

    if len(termino) < 2:

        raise HTTPException(status_code=400, detail="Indica 'consulta' de al menos 2 caracteres para buscar en este tipo.")

    if len(termino) > 120:

        raise HTTPException(status_code=400, detail="La consulta supera 120 caracteres.")

    # Cuando llega solo el consecutivo, anteponemos el codigo del tipo.

    termino_efectivo = categoria.tipo + termino if termino.isdecimal() else termino

    raiz = carpeta_valida(categoria.ruta)

    listado = explorar_pdfs(

        raiz, raiz, termino_efectivo,

        True if recursivo is None else recursivo,

        limite or 100,

    )

    return {

        "clase": categoria.clase,

        "tipo": categoria.tipo,

        "consulta": termino,

        "consulta_efectiva": termino_efectivo,

        "nota_ruta_compartida": resumen_ruta_compartida(categoria),

        **listado,

    }

@app.get("/api/planos/listar", dependencies=[Depends(verificar_clave)], operation_id="listarPlanosTipo")

def listar_planos_tipo(

    tipo: str = Query(...),

    clase: str | None = Query(None),

    carpeta: str = Query("", description="Subcarpeta relativa a la ruta configurada"),

    codigo_pc: str | None = Query(None, description="Obligatorio para PC"),

    recursivo: bool = Query(False),

    limite: int = Query(100, ge=1, le=300),

):

    categoria = seleccionar_tipo(tipo, clase)

    if categoria.tipo == "PC":

        if not codigo_pc:

            raise HTTPException(status_code=400, detail="Indica codigo_pc para listar tipo PC.")

        raiz = carpeta_planos_pc(codigo_pc)

    else:

        raiz = carpeta_valida(categoria.ruta)

        if categoria.ruta == RAICES["TECNICA_2"] and not carpeta.strip():

            raise HTTPException(status_code=400, detail="Para tipos en TECNICA_2 indica una subcarpeta; alternativamente usa /api/planos/buscar.")

    inicio = ruta_segura(raiz, carpeta)

    carpeta_valida(inicio)

    listado = explorar_pdfs(raiz, inicio, "", recursivo, limite)

    return {

        "clase": categoria.clase,

        "tipo": categoria.tipo,

        "carpeta": carpeta or ".",

        "nota_ruta_compartida": resumen_ruta_compartida(categoria),

        **listado,

    }

@app.get("/api/planos/pdf/texto", dependencies=[Depends(verificar_clave)], operation_id="leerPdfPorTipo")

def leer_pdf_por_tipo(

    tipo: str = Query(...),

    archivo: str = Query(..., description="Exactamente 'ruta_en_planos' de buscar o listar"),

    clase: str | None = Query(None),

    codigo_pc: str | None = Query(None),

    max_caracteres: int = Query(50000, ge=1000, le=200000),

):

    categoria, pdf = localizar_pdf(tipo, clase, archivo, codigo_pc)

    return obtener_texto(categoria, pdf, max_caracteres, codigo_pc)

@app.get("/api/planos/pdf/descargar", dependencies=[Depends(verificar_clave)], operation_id="descargarPdfPorTipo")

def descargar_pdf_por_tipo(

    tipo: str = Query(...),

    archivo: str = Query(..., description="Exactamente 'ruta_en_planos' de buscar o listar"),

    clase: str | None = Query(None),

    codigo_pc: str | None = Query(None),

):

    _, pdf = localizar_pdf(tipo, clase, archivo, codigo_pc)

    return FileResponse(str(pdf), filename=pdf.name, media_type="application/pdf")

# Endpoints historicos: se mantienen tal cual en ruta y parametros.

@app.get("/api/planos/{codigo_pc}/pdf/texto", dependencies=[Depends(verificar_clave)], operation_id="leerTextoPdf")

def leer_texto_pdf(

    codigo_pc: str,

    archivo: str = Query(...),

    max_caracteres: int = Query(50000, ge=1000, le=200000),

):

    categoria, pdf = localizar_pdf("PC", None, archivo, codigo_pc)

    return obtener_texto(categoria, pdf, max_caracteres, codigo_pc)

@app.get("/api/planos/{codigo_pc}/pdf/descargar", dependencies=[Depends(verificar_clave)], operation_id="descargarPdf")

def descargar_pdf(codigo_pc: str, archivo: str = Query(...)):

    _, pdf = localizar_pdf("PC", None, archivo, codigo_pc)

    return FileResponse(str(pdf), filename=pdf.name, media_type="application/pdf")

# =========================================================
# SHAREPOINT: REVISION OV (Graph o OneDrive sincronizado)
# =========================================================
import base64
import io
import os
import re
import time
import unicodedata
import zipfile
from pathlib import Path
from urllib.parse import quote, urlparse
import httpx
from fastapi import APIRouter, Depends, HTTPException, Query

GRAPH = 'https://graph.microsoft.com/v1.0'
LECTURA = {'.pdf', '.xlsx', '.xlsm', '.docx', '.pptx', '.txt', '.csv', '.tsv', '.json', '.xml', '.md', '.log'}
IMAGENES = {'.png', '.jpg', '.jpeg', '.webp'}
MAX_BYTES = 30 * 1024 * 1024


def normalizar(s):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', str(s)) if not unicodedata.combining(c)).casefold().split())


def coincide(nombre, consulta):
    n, q = normalizar(nombre), normalizar(consulta)
    # Los consecutivos completos tienen prioridad; 0 no debe escoger 100.
    if q.isdecimal():
        return any(int(x) == int(q) for x in re.findall(r'\d+', n))
    return q in n


def partes_seguras(ruta):
    s = ruta.replace('\\', '/')
    if not s or s.startswith('/') or ':' in s or '\x00' in s:
        raise HTTPException(400, 'Ruta relativa inválida.')
    partes = s.split('/')
    if any(x in {'', '.', '..'} for x in partes):
        raise HTTPException(400, 'Ruta relativa inválida.')
    return partes


class Local:
    modo = 'local_sincronizado'

    def __init__(self):
        ruta = os.getenv('SP_LOCAL_ROOT', '')
        if not ruta:
            raise HTTPException(503, 'Configura SP_LOCAL_ROOT hasta Demetalicos proyectos.')
        self.base = Path(ruta).resolve()
        if not self.base.is_dir():
            raise HTTPException(503, 'La carpeta sincronizada no está disponible para el usuario del servicio.')

    def path(self, ref):
        p = self.base.joinpath(*partes_seguras(ref)) if ref else self.base
        p = p.resolve()
        if p != self.base and self.base not in p.parents:
            raise HTTPException(403, 'Ruta fuera de la raíz autorizada.')
        return p

    def children(self, ref):
        try:
            resultados = []
            for p in self.path(ref).iterdir():
                if p.is_symlink():
                    continue
                self.path(p.relative_to(self.base).as_posix())
                resultados.append({'name': p.name, 'id': p.relative_to(self.base).as_posix(),
                                   'folder': p.is_dir(), 'size': p.stat().st_size if p.is_file() else 0,
                                   'web_url': None})
                if len(resultados) > 10000:
                    raise HTTPException(413, 'Carpeta con más de 10000 elementos; no se devuelve un resultado incompleto.')
            return resultados
        except OSError:
            raise HTTPException(503, 'No se pudo leer la carpeta sincronizada. Verifica permisos y descarga de OneDrive.')

    def content(self, item):
        try:
            with self.path(item['id']).open('rb') as f:
                data = f.read(MAX_BYTES + 1)
        except OSError:
            raise HTTPException(503, 'No se pudo abrir el archivo local. Verifica que esté descargado.')
        if len(data) > MAX_BYTES:
            raise HTTPException(413, 'Archivo mayor de 30 MiB.')
        return data

    def root(self):
        return ''


class Graph:
    modo = 'graph'

    def __init__(self):
        claves = ['SP_TENANT_ID', 'SP_CLIENT_ID', 'SP_CLIENT_SECRET', 'SP_DRIVE_ID']
        faltan = [k for k in claves if not os.getenv(k)]
        if faltan:
            raise HTTPException(503, {'mensaje': 'Falta configuración de Microsoft Graph.', 'variables': faltan})
        self.drive = quote(os.environ['SP_DRIVE_ID'], safe='')
        self.deadline = time.monotonic() + 100
        try:
            r = httpx.post('https://login.microsoftonline.com/' + quote(os.environ['SP_TENANT_ID'], safe='') + '/oauth2/v2.0/token',
                           data={'client_id': os.environ['SP_CLIENT_ID'], 'client_secret': os.environ['SP_CLIENT_SECRET'],
                                 'scope': 'https://graph.microsoft.com/.default', 'grant_type': 'client_credentials'}, timeout=25)
        except httpx.RequestError:
            raise HTTPException(502, 'No fue posible conectar con Microsoft Entra.')
        if r.status_code != 200:
            raise HTTPException(502, 'Microsoft Entra rechazó la autenticación. Revisa tenant, client ID y vigencia del secreto.')
        self.token = r.json()['access_token']

    def get(self, path):
        url = path if path.startswith('https://') else GRAPH + path
        if not url.startswith(GRAPH + '/'):
            raise HTTPException(502, 'Enlace de paginación fuera de Microsoft Graph.')
        if time.monotonic() > self.deadline:
            raise HTTPException(504, 'La consulta superó el tiempo máximo. No se seleccionó una coincidencia parcial.')
        try:
            r = httpx.get(url, headers={'Authorization': 'Bearer ' + self.token}, timeout=25)
        except httpx.RequestError:
            raise HTTPException(502, 'Error de comunicación con Microsoft Graph.')
        if r.status_code >= 400:
            try:
                codigo = r.json().get('error', {}).get('code', 'GraphError')
            except ValueError:
                codigo = 'GraphError'
            raise HTTPException(r.status_code if r.status_code in {403, 404, 429} else 502,
                                {'mensaje': 'Microsoft Graph no pudo completar la consulta.', 'codigo': codigo,
                                 'http_graph': r.status_code, 'retry_after': r.headers.get('Retry-After')})
        return r

    def root(self):
        root = os.getenv('SP_PROJECTS_PATH', 'Demetalicos proyectos').strip('/')
        if root:
            partes_seguras(root)
            item = self.get(f'/drives/{self.drive}/root:/{quote(root, safe="/")}').json()
        else:
            item = self.get(f'/drives/{self.drive}/root').json()
        if 'folder' not in item:
            raise HTTPException(400, 'SP_PROJECTS_PATH debe señalar una carpeta.')
        return item['id']

    def children(self, ref):
        url = f'/drives/{self.drive}/items/{quote(ref, safe="")}/children?$top=200&$select=id,name,size,folder,file,webUrl,remoteItem'
        rows = []
        pages = 0
        while url:
            data = self.get(url).json()
            for x in data.get('value', []):
                if 'remoteItem' in x:
                    continue  # No seguir accesos directos a otro ámbito.
                rows.append({'id': x['id'], 'name': x['name'], 'folder': 'folder' in x,
                             'size': x.get('size', 0), 'web_url': x.get('webUrl')})
            pages += 1
            if len(rows) > 10000 or pages > 100:
                raise HTTPException(413, 'Listado demasiado grande; no se seleccionó una coincidencia parcial.')
            url = data.get('@odata.nextLink')
        return rows

    def content(self, item):
        if item['size'] > MAX_BYTES:
            raise HTTPException(413, 'Archivo mayor de 30 MiB.')
        r = self.get(f'/drives/{self.drive}/items/{quote(item["id"], safe="")}/content')
        if r.status_code in {301, 302, 303, 307, 308}:
            url = r.headers.get('location', '')
            if urlparse(url).scheme != 'https':
                raise HTTPException(502, 'URL de descarga no válida.')
            # URL preautenticada de Microsoft: nunca enviar el bearer a este host.
            try:
                with httpx.stream('GET', url, timeout=45, follow_redirects=True) as response:
                    if response.status_code != 200:
                        raise HTTPException(502, 'Falló la descarga desde SharePoint.')
                    data = bytearray()
                    for chunk in response.iter_bytes():
                        data.extend(chunk)
                        if len(data) > MAX_BYTES:
                            raise HTTPException(413, 'Archivo mayor de 30 MiB.')
                    return bytes(data)
            except httpx.RequestError:
                raise HTTPException(502, 'Se interrumpió la descarga desde SharePoint.')
        if r.status_code != 200:
            raise HTTPException(502, 'Respuesta de descarga inesperada.')
        if len(r.content) > MAX_BYTES:
            raise HTTPException(413, 'Archivo mayor de 30 MiB.')
        return r.content


def proveedor():
    modo = os.getenv('SP_MODE', 'graph').casefold()
    if modo == 'graph':
        return Graph()
    if modo == 'local':
        return Local()
    raise HTTPException(503, 'SP_MODE debe ser graph o local.')


def elegir(items, consulta, nivel, exacta=False):
    if not consulta.strip() or len(consulta) > 180:
        raise HTTPException(400, f'Indica {nivel}, entre 1 y 180 caracteres.')
    carpetas = [x for x in items if x['folder']]
    iguales = [x for x in carpetas if normalizar(x['name']) == normalizar(consulta)]
    matches = iguales or ([] if exacta else [x for x in carpetas if coincide(x['name'], consulta)])
    if not matches:
        raise HTTPException(404, {'mensaje': 'No se encontró carpeta.', 'nivel': nivel, 'consulta': consulta})
    if len(matches) > 1:
        raise HTTPException(409, {'mensaje': 'Varias coincidencias. Repite con el nombre completo de la carpeta.',
                                 'nivel': nivel, 'opciones': [x['name'] for x in matches[:100]],
                                 'total': len(matches)})
    return matches[0]


def resolver(p, cliente, caso, cotizacion, ov):
    ref, nombres = p.root(), []
    for nivel, consulta in [('cliente', cliente), ('caso', caso), ('cotizacion', cotizacion), ('ov', ov),
                            ('revision', '00 - Revision OV')]:
        item = elegir(p.children(ref), consulta, nivel, exacta=nivel == 'revision')
        ref = item['id']
        nombres.append(item['name'])
    return item, '/'.join(nombres)


def localizar_archivo(p, carpeta, archivo):
    partes = partes_seguras(archivo)
    ref = carpeta['id']
    for i, nombre in enumerate(partes):
        candidatos = [x for x in p.children(ref) if x['name'] == nombre]
        if len(candidatos) != 1:
            raise HTTPException(404, 'Archivo no encontrado. Usa ruta_en_revision del listado.')
        item = candidatos[0]
        if i < len(partes) - 1 and not item['folder']:
            raise HTTPException(400, 'Ruta de archivo inválida.')
        ref = item['id']
    if item['folder']:
        raise HTTPException(400, 'La ruta corresponde a una carpeta.')
    return item


def inventario(p, carpeta, recursivo, limite):
    cola = [(carpeta['id'], '', 0)]
    filas, parcial = [], False
    visitados = 0
    while cola:
        ref, prefijo, depth = cola.pop(0)
        for x in p.children(ref):
            visitados += 1
            if visitados > 10000:
                return filas, True
            ruta = prefijo + x['name']
            if x['folder']:
                if recursivo and depth < 8:
                    cola.append((x['id'], ruta + '/', depth + 1))
                elif recursivo:
                    parcial = True
                tipo = 'carpeta'
            else:
                tipo = Path(x['name']).suffix.casefold()
            filas.append({'nombre': x['name'], 'ruta_en_revision': ruta, 'tipo': tipo,
                          'tamano_bytes': x['size'], 'web_url': x['web_url'],
                          'lectura_texto': tipo in LECTURA, 'lectura_visual': tipo in IMAGENES})
            if len(filas) >= limite:
                return filas, True
    return filas, parcial


def fragmentos(data, ext):
    if ext in {'.xlsx', '.xlsm', '.docx', '.pptx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(x.file_size for x in z.infolist()) > 150 * 1024 * 1024:
                raise HTTPException(413, 'Documento descomprimido demasiado grande.')
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        for n, page in enumerate(reader.pages, 1):
            yield f'\n[Página {n}]\n'
            texto = page.extract_text() or ''
            yield texto if texto.strip() else '[Página sin texto extraíble: requiere revisión visual u OCR.]'
    elif ext in {'.xlsx', '.xlsm'}:
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            for sheet in book:
                yield f'\n[Hoja: {sheet.title}]\n'
                # Evita dimensiones declaradas incorrectas y respeta celdas reales del XML.
                sheet.reset_dimensions()
                for n, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if n > 50000:
                        yield '[Límite de 50000 filas por hoja; contenido parcial.]\n'
                        break
                    if any(v is not None for v in row):
                        yield f'Fila {n}: ' + '\t'.join('' if v is None else str(v) for v in row) + '\n'
        finally:
            book.close()
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        for p in doc.paragraphs:
            yield p.text + '\n'
        for n, table in enumerate(doc.tables, 1):
            yield f'\n[Tabla {n}]\n'
            for row in table.rows:
                yield '\t'.join(c.text for c in row.cells) + '\n'
    elif ext == '.pptx':
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        for n, slide in enumerate(prs.slides, 1):
            yield f'\n[Diapositiva {n}]\n'
            for shape in slide.shapes:
                if shape.has_text_frame:
                    yield shape.text + '\n'
                if shape.has_table:
                    for row in shape.table.rows:
                        yield '\t'.join(c.text for c in row.cells) + '\n'
    else:
        try:
            yield data.decode('utf-8-sig')
        except UnicodeDecodeError:
            yield data.decode('cp1252', errors='replace')


def extraer(data, ext, offset, max_caracteres):
    if ext not in LECTURA:
        return {'estado': 'formato_no_soportado', 'texto': '', 'truncado': False,
                'observacion': 'Para PNG/JPG/WEBP usa ver_imagen_revision_ov. Otros formatos requieren conversión.'}
    resultado, total, fin = [], 0, offset + max_caracteres
    mas = False
    gen = fragmentos(data, ext)
    try:
        for bloque in gen:
            anterior = total
            total += len(bloque)
            if total > offset and anterior < fin:
                resultado.append(bloque[max(0, offset - anterior):max(0, fin - anterior)])
            if total > fin:
                mas = True
                break
    finally:
        gen.close()
    texto = ''.join(resultado)
    return {'estado': 'texto_extraido' if texto.strip() else 'sin_texto', 'texto': texto,
            'offset': offset, 'truncado': mas, 'siguiente_offset': fin if mas else None,
            'observacion': 'Extracción de texto; no incluye imágenes incrustadas, OCR ni evaluación de fórmulas. Excel devuelve fórmulas existentes.'}


def crear_router(verificar_clave):
    router = APIRouter(prefix='/api/sharepoint', dependencies=[Depends(verificar_clave)])

    @router.get('/estado')
    def estado_sharepoint():
        p = proveedor()
        p.children(p.root())
        return {'accesible': True, 'modo': p.modo, 'carpeta': os.getenv('SP_PROJECTS_PATH', 'Demetalicos proyectos')}

    @router.get('/revision-ov/buscar')
    def buscar(cliente: str, caso: str, cotizacion: str, ov: str,
               recursivo: bool = False, limite: int = Query(200, ge=1, le=1000)):
        p = proveedor()
        carpeta, ruta = resolver(p, cliente, caso, cotizacion, ov)
        archivos, parcial = inventario(p, carpeta, recursivo, limite)
        return {'modo': p.modo, 'ruta_revision': ruta, 'web_url': carpeta['web_url'],
                'archivos': archivos, 'total_devuelto': len(archivos), 'parcial': parcial,
                'observacion': 'Si parcial=true, aumenta limite hasta 1000 o consulta una OV con menos archivos.' if parcial else None}

    @router.get('/revision-ov/leer')
    def leer(cliente: str, caso: str, cotizacion: str, ov: str, archivo: str,
             offset: int = Query(0, ge=0, le=10000000), max_caracteres: int = Query(30000, ge=1000, le=100000)):
        p = proveedor()
        carpeta, ruta = resolver(p, cliente, caso, cotizacion, ov)
        item = localizar_archivo(p, carpeta, archivo)
        ext = Path(item['name']).suffix.casefold()
        try:
            data = p.content(item) if ext in LECTURA else b''
            result = extraer(data, ext, offset, max_caracteres)
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, 'No se pudo interpretar el archivo. Puede estar dañado, cifrado o tener un formato distinto a su extensión.')
        return {'ruta_revision': ruta, 'archivo': archivo, 'web_url': item['web_url'], **result}

    @router.get('/revision-ov/imagen')
    def imagen(cliente: str, caso: str, cotizacion: str, ov: str, archivo: str):
        p = proveedor()
        carpeta, _ = resolver(p, cliente, caso, cotizacion, ov)
        item = localizar_archivo(p, carpeta, archivo)
        if Path(item['name']).suffix.casefold() not in IMAGENES:
            raise HTTPException(415, 'Solo PNG, JPG, JPEG y WEBP.')
        from PIL import Image, ImageOps
        try:
            with Image.open(io.BytesIO(p.content(item))) as im:
                if im.width * im.height > 25000000:
                    raise HTTPException(413, 'Imagen mayor de 25 megapíxeles.')
                im = ImageOps.exif_transpose(im).convert('RGB')
                im.thumbnail((1600, 1600))
                out = io.BytesIO()
                im.save(out, 'JPEG', quality=85)
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, 'No se pudo interpretar la imagen.')
        return {'archivo': archivo, 'formato': 'jpeg', 'base64': base64.b64encode(out.getvalue()).decode('ascii')}

    return router


app.include_router(crear_router(verificar_clave))

if __name__ == "__main__":

    import uvicorn

    uvicorn.run("app:app", host="127.0.0.1", port=8091, reload=False)
